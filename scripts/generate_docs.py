import os
import re
import sys
import glob
import time
import hashlib
import base64
import urllib.request
from html.parser import HTMLParser
from PIL import Image
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

class StepCardParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.badge = ""
        self.title = ""
        self.items = [] # list of (depth, text)
        self.in_badge = False
        self.in_title = False
        self.ul_depth = 0
        self.li_stack = []

    def handle_starttag(self, tag, attrs):
        attrs_dict = dict(attrs)
        if tag == 'span':
            cls = attrs_dict.get('class', '')
            if 'step-badge' in cls:
                self.in_badge = True
            elif 'step-title' in cls:
                self.in_title = True
        elif tag == 'ul':
            if self.li_stack and self.li_stack[-1]:
                txt = "".join(self.li_stack[-1]).strip()
                if txt:
                    self.items.append((self.ul_depth, txt))
                self.li_stack[-1] = []
            self.ul_depth += 1
        elif tag == 'li':
            self.li_stack.append([])
        elif tag == 'code':
            if self.li_stack:
                self.li_stack[-1].append('<code>')

    def handle_endtag(self, tag):
        if tag == 'span':
            self.in_badge = False
            self.in_title = False
        elif tag == 'ul':
            self.ul_depth = max(0, self.ul_depth - 1)
        elif tag == 'li':
            if self.li_stack:
                accum = self.li_stack.pop()
                txt = "".join(accum).strip()
                if txt:
                    self.items.append((self.ul_depth, txt))
        elif tag == 'code':
            if self.li_stack:
                self.li_stack[-1].append('</code>')

    def handle_data(self, data):
        if self.in_badge:
            self.badge += data
        elif self.in_title:
            self.title += data
        elif self.li_stack:
            self.li_stack[-1].append(data)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top="CCCCCC", bottom="CCCCCC", left="CCCCCC", right="CCCCCC", sz="4", left_sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    def bdr_xml(side, color, size):
        if not color or color.lower() == "none":
            return f'<w:{side} w:val="none"/>'
        return f'<w:{side} w:val="single" w:sz="{size}" w:space="0" w:color="{color}"/>'
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'{bdr_xml("top", top, sz)}'
        f'{bdr_xml("left", left, left_sz)}'
        f'{bdr_xml("bottom", bottom, sz)}'
        f'{bdr_xml("right", right, sz)}'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def add_inline_runs(paragraph, text, default_size=Pt(12), is_heading=False):
    text = re.sub(r'<br\s*/?>', ' ', text)
    text = text.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&").replace("&quot;", '"')
    
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|<code>.*?</code>|<u>.*?</u>)')
    tokens = pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        
        bold = False
        italic = False
        code = False
        underline = False
        content = token
        
        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            bold = True
            content = token[2:-2]
        elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
            italic = True
            content = token[1:-1]
        elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
            code = True
            content = token[1:-1]
        elif token.startswith("<code>") and token.endswith("</code>"):
            code = True
            content = token[6:-7]
        elif token.startswith("<u>") and token.endswith("</u>"):
            underline = True
            content = token[3:-4]
            if content.startswith("**") and content.endswith("**"):
                bold = True
                content = content[2:-2]
        
        if content.startswith("<code>") and content.endswith("</code>"):
            code = True
            content = content[6:-7]
            
        run = paragraph.add_run(content)
        run.bold = True if is_heading else bold
        run.italic = italic
        run.underline = underline
        
        if code:
            run.font.name = 'Consolas'
            run.font.size = Pt(10.5) if not is_heading else default_size
            run.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)
        else:
            run.font.name = 'Times New Roman'
            run.font.size = default_size

def extract_project_title(content):
    sec0 = content.split('\n---\n')[0]
    match = re.search(r'^#\s+(.+)$', sec0, re.MULTILINE)
    if match:
        return match.group(1).strip()
    return "Project"

def render_step_card(doc, card_html):
    parser = StepCardParser()
    parser.feed(card_html)
    
    badge = parser.badge.strip() if parser.badge else "STEP"
    title = parser.title.strip() if parser.title else ""
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(17.8)
    
    cell = table.cell(0, 0)
    set_cell_background(cell, "F4F6F8")
    set_cell_margins(cell, top=60, bottom=60, left=120, right=120)
    set_cell_borders(cell, top="D0D5DD", bottom="D0D5DD", left="2563EB", right="D0D5DD", sz="4", left_sz="24")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    
    header_text = f"{badge} — {title}" if title else badge
    run_badge = p.add_run(header_text)
    run_badge.font.name = 'Times New Roman'
    run_badge.font.size = Pt(11.5)
    run_badge.bold = True
    run_badge.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    for depth, clean_item in parser.items:
        clean_item = clean_item.strip()
        if not clean_item:
            continue
            
        p_item = cell.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(1)
        p_item.paragraph_format.line_spacing = 1.05
        
        indent_cm = 0.4 * depth
        p_item.paragraph_format.left_indent = Cm(indent_cm)
        
        if clean_item.startswith("[ ]"):
            clean_item = "☐  " + clean_item[3:].strip()
            item_font_size = Pt(11)
        elif clean_item.startswith("[x]"):
            clean_item = "☑  " + clean_item[3:].strip()
            item_font_size = Pt(11)
        elif depth > 1:
            clean_item = "– " + clean_item
            item_font_size = Pt(10.5)
        else:
            clean_item = "•  " + clean_item
            item_font_size = Pt(11)
            
        add_inline_runs(p_item, clean_item, default_size=item_font_size)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(2)
    p_spacer.paragraph_format.line_spacing = Pt(1)

def render_table(doc, table_lines):
    raw_rows = []
    for line in table_lines:
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        raw_rows.append(cells)
    
    if len(raw_rows) < 2:
        return
    
    header = raw_rows[0]
    is_sep = any(set(c).issubset({'-', ':', ' '}) for c in raw_rows[1])
    data_rows = raw_rows[2:] if is_sep else raw_rows[1:]
    
    num_cols = len(header)
    num_rows = 1 + len(data_rows)
    
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    
    col_width = Cm(17.8 / num_cols)
    for col in table.columns:
        col.width = col_width
    
    for col_idx, text in enumerate(header):
        cell = table.cell(0, col_idx)
        set_cell_background(cell, "EAECEF")
        set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
        set_cell_borders(cell, top="CCCCCC", bottom="999999", left="CCCCCC", right="CCCCCC", sz="4")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        add_inline_runs(p, text, default_size=Pt(10.5), is_heading=True)
    
    for row_idx, rdata in enumerate(data_rows):
        bg = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx in range(num_cols):
            cell = table.cell(row_idx + 1, col_idx)
            text = rdata[col_idx] if col_idx < len(rdata) else ""
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=30, bottom=30, left=60, right=60)
            set_cell_borders(cell, top="E5E7EB", bottom="E5E7EB", left="E5E7EB", right="E5E7EB", sz="4")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            add_inline_runs(p, text, default_size=Pt(10.5))
            
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(2)
    p_spacer.paragraph_format.line_spacing = Pt(1)

def render_code_block(doc, code_lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(17.8)
    
    cell = table.cell(0, 0)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
    set_cell_borders(cell, top="E2E8F0", bottom="E2E8F0", left="CBD5E1", right="E2E8F0", sz="4", left_sz="12")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    
    full_code = "\n".join(code_lines)
    run = p.add_run(full_code)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(2)
    p_spacer.paragraph_format.line_spacing = Pt(1)

def render_mermaid_diagram(doc, mermaid_code):
    clean = mermaid_code.strip()
    h = hashlib.md5(clean.encode('utf-8')).hexdigest()
    cache_dir = os.path.join(os.getcwd(), 'docs', 'diagrams')
    os.makedirs(cache_dir, exist_ok=True)
    img_path = os.path.join(cache_dir, f"{h}.png")
    
    if not os.path.exists(img_path):
        sanitized = clean.replace('>=', '&ge;').replace('Play Again? (y/n)', 'Play Again (y/n)?')
        for variant in [clean, sanitized]:
            downloaded = False
            for attempt in range(3):
                try:
                    enc = base64.b64encode(variant.encode('utf-8')).decode('ascii')
                    url = f"https://mermaid.ink/img/{enc}"
                    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
                    with urllib.request.urlopen(req, timeout=12) as r:
                        data = r.read()
                        if len(data) > 100:
                            with open(img_path, 'wb') as f:
                                f.write(data)
                            downloaded = True
                            break
                except Exception:
                    time.sleep(1.5)
            if downloaded:
                break
                
    if os.path.exists(img_path):
        try:
            with Image.open(img_path) as img:
                w_px, h_px = img.size
                
            aspect = w_px / h_px
            target_w = min(15.0, 11.5 * aspect)
            target_h = target_w / aspect
            if target_h > 11.5:
                target_h = 11.5
                target_w = target_h * aspect
                
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run()
            run.add_picture(img_path, width=Cm(target_w), height=Cm(target_h))
            return
        except Exception as e:
            print(f"Warning: Failed to insert picture {img_path}: {e}")
            
    render_code_block(doc, mermaid_code.splitlines())

def render_section_content(doc, section_text):
    lines = section_text.split('\n')
    idx = 0
    total_lines = len(lines)
    
    while idx < total_lines:
        line = lines[idx]
        stripped = line.strip()
        
        if not stripped:
            idx += 1
            continue
            
        # 1. Step Card check
        if '<div class="step-card">' in line:
            card_lines = [line]
            idx += 1
            nested_divs = 1
            while idx < total_lines:
                c_line = lines[idx]
                card_lines.append(c_line)
                if '<div' in c_line:
                    nested_divs += c_line.count('<div')
                if '</div>' in c_line:
                    nested_divs -= c_line.count('</div>')
                idx += 1
                if nested_divs <= 0:
                    break
            render_step_card(doc, "\n".join(card_lines))
            continue
            
        # 2. Code Block / Mermaid check
        if stripped.startswith("```"):
            is_mermaid = "mermaid" in stripped.lower()
            code_lines = []
            idx += 1
            while idx < total_lines and not lines[idx].strip().startswith("```"):
                code_lines.append(lines[idx])
                idx += 1
            if idx < total_lines:
                idx += 1
                
            if is_mermaid:
                render_mermaid_diagram(doc, "\n".join(code_lines))
            else:
                render_code_block(doc, code_lines)
            continue
            
        # 3. Table check
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = [stripped]
            idx += 1
            while idx < total_lines and lines[idx].strip().startswith("|") and lines[idx].strip().endswith("|"):
                table_lines.append(lines[idx].strip())
                idx += 1
            render_table(doc, table_lines)
            continue
            
        # 4. Headings
        if stripped.startswith("#"):
            heading_match = re.match(r'^(#+)\s+(.+)$', stripped)
            if heading_match:
                level = len(heading_match.group(1))
                heading_text = heading_match.group(2).strip()
                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                add_inline_runs(p, heading_text, default_size=Pt(14), is_heading=True)
                idx += 1
                continue
                
        # 5. Milestone Checkpoint or Subheading like **Title** on its own line
        concept_match = re.match(r'^\*\*([^*]+)\*\*$', stripped)
        if concept_match:
            concept_text = concept_match.group(1).strip()
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            add_inline_runs(p, concept_text, default_size=Pt(13), is_heading=True)
            idx += 1
            continue
            
        # 6. Checkbox lists: - [ ] or - [x]
        if stripped.startswith("- [ ]") or stripped.startswith("- [x]"):
            is_checked = stripped.startswith("- [x]")
            box_sym = "☑  " if is_checked else "☐  "
            item_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.08
            p.paragraph_format.left_indent = Cm(0.5)
            p.add_run(box_sym).font.name = 'Times New Roman'
            add_inline_runs(p, item_text, default_size=Pt(12))
            idx += 1
            continue
            
        # 7. Bullet lists: - or *
        if (stripped.startswith("- ") or stripped.startswith("* ")) and not stripped.startswith("**"):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.08
            p.paragraph_format.left_indent = Cm(0.5)
            p.add_run("•  ").font.name = 'Times New Roman'
            add_inline_runs(p, bullet_text, default_size=Pt(12))
            idx += 1
            continue
            
        # 8. Numbered lists: 1. , 2. 
        num_match = re.match(r'^(\d+)\.\s+(.+)$', stripped)
        if num_match:
            num = num_match.group(1)
            item_text = num_match.group(2).strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.08
            p.paragraph_format.left_indent = Cm(0.5)
            p.add_run(f"{num}.  ").font.name = 'Times New Roman'
            add_inline_runs(p, item_text, default_size=Pt(12))
            idx += 1
            continue
            
        # 9. Blockquote: > text
        if stripped.startswith(">"):
            bq_text = stripped.lstrip("> ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.08
            p.paragraph_format.left_indent = Cm(0.6)
            add_inline_runs(p, bq_text, default_size=Pt(11.5))
            idx += 1
            continue

        # 10. Display Math: $$ ... $$
        if stripped.startswith("$$") and stripped.endswith("$$") and len(stripped) >= 4:
            math_text = stripped[2:-2].strip()
            math_clean = re.sub(r'\\text\{([^}]+)\}', r'\1', math_text)
            math_clean = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1) / (\2)', math_clean)
            math_clean = math_clean.replace(r'\sum_{i=1}^n', '∑(i=1 to n)')
            math_clean = math_clean.replace(r'\sum', '∑')
            math_clean = math_clean.replace(r'\dots', '...')
            math_clean = math_clean.replace(r'\hat{y}_i', 'ŷ_i')
            math_clean = math_clean.replace(r'\times', '×')
            math_clean = math_clean.replace(r'\%', '%')
            math_clean = math_clean.replace('\\', '')
            
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(math_clean)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.italic = True
            run.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            idx += 1
            continue

        # 11. Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.08
        add_inline_runs(p, stripped, default_size=Pt(12))
        idx += 1

def convert_markdown_to_docx(md_path, output_docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
        
    project_title = extract_project_title(content)
    print(f"Converting '{md_path}' -> '{output_docx_path}' (Title: {project_title})")
    
    sections = re.split(r'\n\s*---\s*\n', content)
    print(f"  Found {len(sections)} sections in markdown.")
    
    doc = docx.Document()
    
    # Margins 1.6 cm on all sides
    for s in doc.sections:
        s.top_margin = Cm(1.6)
        s.bottom_margin = Cm(1.6)
        s.left_margin = Cm(1.6)
        s.right_margin = Cm(1.6)
        
    # Normal style: Times New Roman, 12 pt
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    
    # Skip Section 0 (metadata banner) and Section 1 (Curriculum Outline)
    curriculum_sections = sections[2:]
    
    # 1. Split Section 2 (Chapter 1) into 3 distinct pages
    sec2_text = curriculum_sections[0].strip()
    m_learn = re.search(r'\*\*What You Will Learn\*\*', sec2_text)
    m_need = re.search(r'\*\*What You Need\*\*', sec2_text)
    
    if m_learn and m_need:
        part1 = sec2_text[:m_learn.start()].strip()
        part2 = sec2_text[m_learn.start():m_need.start()].strip()
        part3 = sec2_text[m_need.start():].strip()
        
        part2 = re.sub(r'^\*\*What You Will Learn\*\*', '## What You Will Learn', part2)
        part3 = re.sub(r'^\*\*What You Need\*\*', '## What You Need', part3)
        
        all_pages = [part1, part2, part3]
    else:
        all_pages = [sec2_text]
        
    # 2. Process remaining sections, splitting Chronological Flow and Thinking in OOP
    remaining_sections = curriculum_sections[1:]
    for sec_raw in remaining_sections:
        sec_clean = sec_raw.strip()
        if not sec_clean:
            continue
            
        sec_lower = sec_clean.lower()
        has_oop = 'thinking in oop' in sec_lower
        has_chrono = 'chronological' in sec_lower or 'game flow' in sec_lower
        
        if has_oop and has_chrono:
            m_oop = re.search(r'\*\*Thinking in OOP[^\n]*\*\*', sec_clean, re.IGNORECASE)
            if m_oop:
                part_chrono = sec_clean[:m_oop.start()].strip()
                part_oop = sec_clean[m_oop.start():].strip()
                # Promote Thinking in OOP to a 14pt heading
                part_oop = re.sub(r'^\*\*Thinking in OOP([^*]*)\*\*', r'### 2.3 System Architecture: Thinking in OOP\1', part_oop, flags=re.IGNORECASE)
                all_pages.append(part_chrono)
                all_pages.append(part_oop)
                continue
                
        # Split 2.2 System Architecture if co-located with 2.1
        m_arch = re.search(r'\n(###\s*2\.2\s+System Architecture[^\n]*)', sec_clean)
        if m_arch:
            part_intro = sec_clean[:m_arch.start()].strip()
            part_arch = sec_clean[m_arch.start():].strip()
            all_pages.append(part_intro)
            all_pages.append(part_arch)
            continue

        all_pages.append(sec_clean)
        
    print(f"  Total pages to render: {len(all_pages)} (with intro and OOP splits applied).")
    
    for page_idx, page_text in enumerate(all_pages):
        cleaned_text = page_text.strip()
        if not cleaned_text:
            continue
            
        if page_idx > 0:
            doc.add_page_break()
            
        if page_idx == 0:
            # Page 1: Only the Project Name at 16pt Bold on top!
            p_title = doc.add_paragraph()
            p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_title.paragraph_format.space_before = Pt(0)
            p_title.paragraph_format.space_after = Pt(6)
            run_title = p_title.add_run(project_title)
            run_title.font.name = 'Times New Roman'
            run_title.font.size = Pt(16)
            run_title.bold = True
            
        render_section_content(doc, cleaned_text)
        
    doc.save(output_docx_path)
    print(f"  Successfully saved '{output_docx_path}'!")

def main():
    docs_dir = os.path.join(os.getcwd(), 'docs')
    os.makedirs(docs_dir, exist_ok=True)
    
    projects = [
        ("Project 1 - Bingo Game - I/BingoGame.md", "docs/BingoGame.docx"),
        ("Project 2 - Stone Paper Scissors Game - B/StonePaperScissors.md", "docs/StonePaperScissors.docx"),
        ("Project 3 -  Hangman - B/Hangman.md", "docs/Hangman.docx"),
        ("Project 4 - Blackjack Game - B/Blackjack.md", "docs/Blackjack.docx"),
        ("Project 5 - Multi Unit & Currency Converter - B/Converter.md", "docs/Converter.docx"),
        ("Project 6 - AI Joke & Riddle Generator - B/JokeRiddleGenerator.md", "docs/JokeRiddleGenerator.docx"),
        ("Project 7 - Credit Card Fraud Detection - B/CreditCardFraud.md", "docs/CreditCardFraud.docx"),
        ("Project 8 - House Price Predictor - B/HousePricePredictor.md", "docs/HousePricePredictor.docx"),
        ("Project 9 - Cats vs Dogs  - Binary Image (Deep Learning)/CatsVsDogs.md", "docs/CatsVsDogs.docx"),
        ("Project 10 - Customer Churn & Retention Predictor - I/CustomerChurn.md", "docs/CustomerChurn.docx"),
    ]
    
    # Optional CLI filtering: python generate_docs.py 6 7 8 9 10 or python generate_docs.py --ai
    targets = sys.argv[1:]
    if targets:
        if "--ai" in targets:
            selected = [p for p in projects if any(f"Project {i} " in p[0] for i in range(6, 11))]
        else:
            selected = []
            for p in projects:
                for t in targets:
                    if f"Project {t} " in p[0] or t.lower() in p[1].lower():
                        selected.append(p)
                        break
        if selected:
            projects = selected
    
    for md_rel, out_rel in projects:
        md_path = os.path.abspath(md_rel)
        out_path = os.path.abspath(out_rel)
        if os.path.exists(md_path):
            convert_markdown_to_docx(md_path, out_path)
        else:
            print(f"Warning: File not found: {md_path}")

if __name__ == "__main__":
    main()
